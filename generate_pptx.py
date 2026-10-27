"""
Script to generate a college-ready, innovative 16:9 PowerPoint presentation (.pptx)
for the Virtual Daycare Platform project.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    COLOR_BG_DARK = RGBColor(15, 23, 42)      # Deep Slate #0F172A
    COLOR_CARD_DARK = RGBColor(30, 41, 59)    # Slate 800 #1E293B
    COLOR_CARD_BORDER = RGBColor(51, 65, 85) # Slate 700
    COLOR_PRIMARY = RGBColor(99, 102, 241)   # Indigo 500 #6366F1
    COLOR_ACCENT = RGBColor(16, 185, 129)    # Emerald 500 #10B981
    COLOR_TEXT_LIGHT = RGBColor(248, 250, 252)# Slate 50
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184)# Slate 400
    COLOR_AMBER = RGBColor(245, 158, 11)     # Amber 500
    COLOR_CYAN = RGBColor(6, 182, 212)       # Cyan 500
    COLOR_ROSE = RGBColor(244, 63, 94)       # Rose 500

    def add_header(slide, title_text, category_text="VIRTUAL DAYCARE PLATFORM"):
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category_text.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_ACCENT

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_LIGHT

        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.04)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_PRIMARY
        line.line.color.rgb = COLOR_PRIMARY

    def set_slide_background(slide, color=COLOR_BG_DARK):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_DARK, border_color=COLOR_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    # -------------------------------------------------------------
    # SLIDE 1: Title & Capstone Overview
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, COLOR_BG_DARK)

    hero = add_card(slide1, Inches(1.0), Inches(1.0), Inches(11.33), Inches(5.5))
    hero_tf = hero.text_frame
    hero_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    hero_tf.word_wrap = True

    p = hero_tf.paragraphs[0]
    p.text = "🎈 VIRTUAL DAYCARE PLATFORM"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_LIGHT

    p2 = hero_tf.add_paragraph()
    p2.text = "A Full-Stack, Real-Time Early Childhood Management & Parental Engagement Ecosystem"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(18)
    p2.font.color.rgb = COLOR_ACCENT

    p3 = hero_tf.add_paragraph()
    p3.text = "\nProduction Architecture • WebSockets • Daily Attendance Roster • RBAC Privacy Isolation"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_TEXT_MUTED

    p4 = hero_tf.add_paragraph()
    p4.text = "\n\nPresented for: College Capstone Project Review & Viva Defense"
    p4.alignment = PP_ALIGN.CENTER
    p4.font.size = Pt(13)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_CYAN

    slide1.notes_slide.notes_text_frame.text = (
        "Welcome professors and evaluators. Today, I am proud to present the Virtual Daycare Platform, "
        "a production-ready, full-stack web application engineered with Python, Flask, SQLAlchemy, "
        "and WebSockets. It bridges the critical communication and safety gap between early childhood "
        "caretakers and working parents through real-time telemetry, streamlined daily attendance, and strict data isolation."
    )

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement & Motivation
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, COLOR_BG_DARK)
    add_header(slide2, "Problem Statement & Industry Motivation", "1. RESEARCH & PROBLEM DOMAIN")

    col_w = Inches(3.64)
    col_gap = Inches(0.4)
    start_left = Inches(0.8)
    card_top = Inches(1.8)
    card_h = Inches(5.0)

    # Card 1: Communication Blackout
    c1 = add_card(slide2, start_left, card_top, col_w, card_h)
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "🚨 The Communication Gap\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ROSE

    items1 = [
        "Working parents face 8-10 hours of informational blackout regarding their infants and toddlers.",
        "Unorganized chat groups (e.g. WhatsApp) compromise family privacy and cause message overload.",
        "Crucial updates on nutrition, naps, medications, and diaper changes are frequently delayed or forgotten."
    ]
    for it in items1:
        pi = tf1.add_paragraph()
        pi.text = f"• {it}\n"
        pi.font.size = Pt(13)
        pi.font.color.rgb = COLOR_TEXT_LIGHT

    # Card 2: Manual Paper Inefficiencies
    c2 = add_card(slide2, start_left + col_w + col_gap, card_top, col_w, card_h)
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "📋 Manual Paper Logs\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER

    items2 = [
        "Traditional daycares rely on paper sign-in registers, susceptible to fraud, lost sheets, and time-theft.",
        "Zero automated percentage tracking for monthly attendance and regulatory childcare compliance.",
        "Handwritten incident logs are difficult to audit, trace, or share during medical emergencies."
    ]
    for it in items2:
        pi = tf2.add_paragraph()
        pi.text = f"• {it}\n"
        pi.font.size = Pt(13)
        pi.font.color.rgb = COLOR_TEXT_LIGHT

    # Card 3: Security & Privacy Leaks
    c3 = add_card(slide2, start_left + (col_w + col_gap)*2, card_top, col_w, card_h)
    tf3 = c3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "🔒 Privacy & Safety Risks\n"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    items3 = [
        "Child photos and sensitive allergy data are shared indiscriminately in public messenger threads.",
        "No verified identity check during student drop-off and pickup by guardians.",
        "Unvetted educators or staff can access classroom data without centralized admin verification."
    ]
    for it in items3:
        pi = tf3.add_paragraph()
        pi.text = f"• {it}\n"
        pi.font.size = Pt(13)
        pi.font.color.rgb = COLOR_TEXT_LIGHT

    slide2.notes_slide.notes_text_frame.text = (
        "Here we address why this system was conceived. In traditional early childhood centers, parents face anxiety "
        "from lack of visibility, staff are overwhelmed by paper logs, and student privacy is compromised when photos "
        "and medical data are posted on generic messaging platforms. Our solution provides an enterprise-grade, "
        "secure architecture tailored specifically for early childhood education."
    )

    # -------------------------------------------------------------
    # SLIDE 3: Proposed Solution & Core Value
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, COLOR_BG_DARK)
    add_header(slide3, "The Proposed Solution: Virtual Daycare Ecosystem", "2. ARCHITECTURAL OVERVIEW")

    q_w = Inches(5.6)
    q_h = Inches(2.35)
    row1_top = Inches(1.8)
    row2_top = Inches(4.5)
    left1 = Inches(0.8)
    left2 = Inches(6.9)

    quads = [
        ("📱 Real-Time Activity Telemetry", "Instant digital logging for meals, nap cycles, moods, diaper changes, and developmental milestones with instant parent timeline sync.", COLOR_ACCENT, left1, row1_top),
        ("⚡ Daily Attendance Roster & Analytics", "Streamlined 1-click arrival check-in and pickup check-out with automated monthly percentage analytics.", COLOR_PRIMARY, left2, row1_top),
        ("🛡️ Strict Multi-Tenant Child Isolation", "Enforced RBAC decorators and object-level child authorization ensuring zero cross-tenant data visibility (403 Forbidden enforcement).", COLOR_ROSE, left1, row2_top),
        ("🚨 Emergency Alert & Learning Hub", "High-priority medical/accident alerts with external SMS/Email notification hooks and an age-curated interactive early learning library.", COLOR_CYAN, left2, row2_top),
    ]

    for title, desc, color, l, t in quads:
        card = add_card(slide3, l, t, q_w, q_h)
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = color
        
        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_TEXT_LIGHT

    slide3.notes_slide.notes_text_frame.text = (
        "Our platform replaces disjointed manual processes with four cohesive pillars: real-time telemetry "
        "of child routines, streamlined daily attendance tracking, enterprise-grade multi-tenant privacy enforcement, "
        "and critical incident alerting combined with an early learning curriculum."
    )

    # -------------------------------------------------------------
    # SLIDE 4: Full-Stack System Architecture
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, COLOR_BG_DARK)
    add_header(slide4, "Multi-Tier Software Architecture", "3. SYSTEM DESIGN")

    layer_w = Inches(11.7)
    layer_h = Inches(1.1)
    layer_gap = Inches(0.2)
    start_y = Inches(1.8)

    layers = [
        ("Layer 1: Client & Presentation Tier", "Responsive HTML5, Jinja2 Templates, Bootstrap 5, Bootstrap Icons, SocketIO Client, AJAX Hooks", COLOR_PRIMARY),
        ("Layer 2: Application & Real-Time Gateway", "Flask 3.x Application Factory, WSGI Entry, Flask-SocketIO (WebSockets / Polling), Flask-Login Session Management", COLOR_ACCENT),
        ("Layer 3: Modular Domain Blueprints", "Auth, Admin (Director), Caretaker (Roster), Parent (Timeline), Messages (Chat), Learning Hub, REST API (/api)", COLOR_CYAN),
        ("Layer 4: Data & Security Foundation", "SQLAlchemy ORM, SQLite / PostgreSQL, Werkzeug Security, Notification Hooks, CSRF Protection", COLOR_AMBER)
    ]

    for i, (title, details, color) in enumerate(layers):
        y = start_y + i * (layer_h + layer_gap)
        card = add_card(slide4, Inches(0.8), y, layer_w, layer_h)
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = color

        p2 = tf.add_paragraph()
        p2.text = details
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_TEXT_LIGHT

    slide4.notes_slide.notes_text_frame.text = (
        "This diagram outlines our 4-tier architectural blueprint. Layer 1 is the responsive frontend rendered through Jinja2 and Bootstrap 5. "
        "Layer 2 handles request routing and real-time duplex communication using Flask-SocketIO. Layer 3 segregates business logic into 7 specialized "
        "Blueprints for maintainability and scalability. Layer 4 manages persistence, encryption, and device integration."
    )

    # -------------------------------------------------------------
    # SLIDE 5: Role-Based Access Control (RBAC)
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, COLOR_BG_DARK)
    add_header(slide5, "Role-Based Access Control (RBAC) & Security Gatekeeper", "4. SECURITY & USER ROLES")

    card_w = Inches(3.64)
    start_left = Inches(0.8)
    card_top = Inches(1.8)
    card_h = Inches(5.0)

    roles = [
        ("👑 Daycare Director (Admin)", COLOR_AMBER, [
            "Full operational supervision & center analytics.",
            "Caretaker verification workflow (approves or denies new educator accounts).",
            "Child enrollment & caretaker assignment.",
            "Curriculum authoring in Learning Hub.",
            "Parent complaint and inquiry resolution."
        ]),
        ("👩‍🏫 Educator / Caretaker", COLOR_ACCENT, [
            "Gated access: blocked until Director approval.",
            "Access limited strictly to assigned classroom roster.",
            "Streamlined daily attendance check-in & check-out roster.",
            "Multi-category activity logging (meals, naps, etc.).",
            "Broadcast high-priority emergency alerts."
        ]),
        ("👨‍👩‍👧 Parent / Guardian", COLOR_PRIMARY, [
            "Strict tenant isolation: access only to own children.",
            "Real-time synchronized activity timeline.",
            "Private child photo gallery & historical filters.",
            "Attendance history and monthly compliance tracking.",
            "Two-way private chat with assigned caretaker."
        ])
    ]

    for i, (role_title, color, perms) in enumerate(roles):
        x = start_left + i * (card_w + Inches(0.4))
        card = add_card(slide5, x, card_top, card_w, card_h)
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{role_title}\n"
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = color

        for perm in perms:
            pi = tf.add_paragraph()
            pi.text = f"✔ {perm}\n"
            pi.font.size = Pt(12)
            pi.font.color.rgb = COLOR_TEXT_LIGHT

    slide5.notes_slide.notes_text_frame.text = (
        "Security is paramount in child care. We implement a strict three-tier RBAC model. The Director acts as the gatekeeper, "
        "approving newly registered staff before they can view any child records. Parents have zero visibility into any other children, "
        "guaranteed by our custom security decorators and database filters."
    )

    # -------------------------------------------------------------
    # SLIDE 6: Daily Attendance Roster & Compliance Engine
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, COLOR_BG_DARK)
    add_header(slide6, "Daily Attendance Roster & Compliance Engine", "5. CORE INNOVATION 1")

    c_left = add_card(slide6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf_l = c_left.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "⚙️ 1-Click Roster Check-In & Check-Out Flow"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT

    flow_steps = [
        "1. Classroom Roster View: Caretakers access an organized daily view of all assigned children.",
        "2. 1-Click Check-In: Instant single-click arrival mark logs precise UTC timestamp and marks child Present.",
        "3. Evening Check-Out: Pick-up recorded with single click, calculating exact daycare duration.",
        "4. Absence & Custom Notes: Record verified absence or arrival notes with complete staff accountability.",
        "5. Automated Notification: Parents receive real-time push notification the moment attendance is marked."
    ]
    for step in flow_steps:
        pi = tf_l.add_paragraph()
        pi.text = f"\n{step}"
        pi.font.size = Pt(12)
        pi.font.color.rgb = COLOR_TEXT_LIGHT

    c_right = add_card(slide6, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.0))
    tf_r = c_right.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "📊 Automated Analytics & Value"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    value_items = [
        "Eliminates Manual Inefficiencies: Zero paper sign-in queues, no handwriting legibility errors.",
        "Audit-Proof Timestamps: Precise UTC check-in and check-out logs stored securely in the database.",
        "Monthly Percentage Computation: Dynamically counts actual weekdays in the month vs. attended days via 'get_monthly_attendance_stats()'.",
        "Instant Parent Notifications: SocketIO alerts parents the exact second their child enters or leaves the center."
    ]
    for v in value_items:
        pi = tf_r.add_paragraph()
        pi.text = f"\n✔ {v}"
        pi.font.size = Pt(12)
        pi.font.color.rgb = COLOR_TEXT_LIGHT

    slide6.notes_slide.notes_text_frame.text = (
        "Here is our core daily attendance engine: a streamlined daily roster where caretakers mark child arrival and departure with a single click. "
        "When a caretaker marks a child present, the state machine logs arrival timestamp and notifies the parents immediately. "
        "In the evening, pickup marks departure and calculates hours in center, automatically updating monthly compliance analytics."
    )

    # -------------------------------------------------------------
    # SLIDE 7: Real-Time Activity Logging
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, COLOR_BG_DARK)
    add_header(slide7, "Real-Time Activity Telemetry & Daily Logging", "6. CORE INNOVATION 2")

    categories = [
        ("🍼 Nutrition & Meals", "Breakfast, Lunch, Snacks, Fluid Intake, Portions consumed (All, Half, Picky)", COLOR_AMBER),
        ("😴 Nap & Sleep Cycles", "Sleep start time, wake time, duration calculation, sleep quality rating", COLOR_PRIMARY),
        ("😊 Mood & Behavior", "Cheerful, Playful, Calm, Sensitive, Fussy with teacher contextual observations", COLOR_ACCENT),
        ("🧷 Diaper & Hygiene", "Wet, Dirty, Dry diaper changes, Potty training milestones, barrier cream notes", COLOR_CYAN),
        ("🌟 Developmental Milestones", "First steps, social sharing, verbal advances, creative artwork achievements", COLOR_ROSE),
        ("📝 General Educator Notes", "Special requests, clothing changes, medication administration updates", COLOR_TEXT_LIGHT),
    ]

    for i, (cat, desc, col) in enumerate(categories):
        col_idx = i % 3
        row_idx = i // 3
        x = Inches(0.8) + col_idx * Inches(3.95)
        y = Inches(1.8) + row_idx * Inches(2.45)
        card = add_card(slide7, x, y, Inches(3.7), Inches(2.25))
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = cat
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = COLOR_TEXT_LIGHT

    slide7.notes_slide.notes_text_frame.text = (
        "Caregivers log routines through an optimized touch-friendly interface. There are 6 specialized activity categories, "
        "each mapped to distinct Bootstrap UI icons and color tokens. When logged, records are stamped with author IDs, child IDs, "
        "and timestamps, immediately synchronizing with the parent's live timeline."
    )

    # -------------------------------------------------------------
    # SLIDE 8: Real-Time WebSockets & Emergency Alerts
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, COLOR_BG_DARK)
    add_header(slide8, "Real-Time WebSockets & Emergency Broadcast Dispatch", "7. CORE INNOVATION 3")

    c_left = add_card(slide8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf_l = c_left.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "⚡ WebSocket Event Infrastructure"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    ws_items = [
        "Powered by Flask-SocketIO with Socket.IO client library.",
        "Targeted User Rooms: 'user_{user_id}' enables instantaneous delivery of unread notification bell increments.",
        "Scoped Child Chat Rooms: 'child_{child_id}' isolates two-way parent-caregiver messaging threads.",
        "Resilient Fallback: Built-in AJAX polling endpoints (/api/notifications/unread and /api/messages/<id>) ensure zero failure on restricted networks."
    ]
    for it in ws_items:
        pi = tf_l.add_paragraph()
        pi.text = f"\n• {it}"
        pi.font.size = Pt(12)
        pi.font.color.rgb = COLOR_TEXT_LIGHT

    c_right = add_card(slide8, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.0))
    tf_r = c_right.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "🚨 Critical Incident Response System"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ROSE

    alert_items = [
        "Priority Alert Flagging: Educators can trigger instant emergency alerts for fevers, accidents, or urgent medication needs.",
        "Persistent Warning Banners: Emits a high-visibility warning banner directly across the parent dashboard.",
        "Multi-Channel Dispatch Hooks: Integrated with 'notifications.py' featuring ready-to-wire Twilio SMS and SendGrid Email dispatchers.",
        "Immediate Actionability: Includes 1-click caregiver calling and medical advisory acknowledgments."
    ]
    for it in alert_items:
        pi = tf_r.add_paragraph()
        pi.text = f"\n• {it}"
        pi.font.size = Pt(12)
        pi.font.color.rgb = COLOR_TEXT_LIGHT

    slide8.notes_slide.notes_text_frame.text = (
        "For communication, we use Flask-SocketIO to create scoped rooms. A user room for notification bell updates, "
        "and a child room for instant chat. If a child develops a fever or has an accident, the caregiver can trigger an emergency alert. "
        "This bypasses ordinary chat, paints a high-priority advisory banner on the parent screen, and triggers external SMS/email hooks."
    )

    # -------------------------------------------------------------
    # SLIDE 9: Interactive Early Learning Hub
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, COLOR_BG_DARK)
    add_header(slide9, "Interactive Early Childhood Learning Hub", "8. CURRICULUM MODULE")

    hub_cards = [
        ("👶 Infant Track (0 - 1 yr)", "Sensory tracking activities, lullabies, high-contrast visual cues, and gentle motor movement exercises designed for early neural development.", COLOR_CYAN, Inches(0.8)),
        ("🧸 Toddler Track (1 - 3 yrs)", "Interactive rhymes (Itsy Bitsy Spider, Wheels on the Bus), tactile color sorting, animal sound recognition, and cooperative play.", COLOR_AMBER, Inches(4.75)),
        ("🎨 Preschool Track (3 - 5 yrs)", "Phonics storytelling, basic counting, social-emotional learning, kindergarten readiness modules, and fine-motor craft guides.", COLOR_ACCENT, Inches(8.7))
    ]

    for title, desc, col, x in hub_cards:
        card = add_card(slide9, x, Inches(1.8), Inches(3.8), Inches(5.0))
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{title}\n"
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}\n\n"
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_TEXT_LIGHT

        p3 = tf.add_paragraph()
        p3.text = "Features:\n• Immersive Reading Mode\n• Audio/Video Embed Support\n• Director Curriculum CMS\n• Age-Filtered Search"
        p3.font.size = Pt(12)
        p3.font.color.rgb = COLOR_TEXT_MUTED

    slide9.notes_slide.notes_text_frame.text = (
        "Beyond daycare operations, the platform integrates an Early Learning Hub. It categorizes nursery rhymes, "
        "interactive stories, and sensory activities specifically tailored to developmental age groups: Infants, Toddlers, and Preschoolers. "
        "Parents can continue classroom learning at home in an ad-free, secure environment."
    )

    # -------------------------------------------------------------
    # SLIDE 10: Database Schema & Entity Relationships
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, COLOR_BG_DARK)
    add_header(slide10, "Relational Database Schema & Data Models", "9. DATA ARCHITECTURE")

    col_w = Inches(3.64)
    start_left = Inches(0.8)

    entities = [
        ("User & Child Entities", COLOR_PRIMARY, [
            "User: id, email, password_hash, name, role ('admin'|'caregiver'|'parent'), is_approved, phone.",
            "Child: id, name, dob, gender, allergies, notes, photo_url, qr_code_path.",
            "Relationships: Foreign keys linking parent_id and caregiver_id with cascade deletion."
        ]),
        ("Activity & Attendance", COLOR_ACCENT, [
            "ActivityLog: id, child_id, type ('meal'|'nap'|'mood'|'diaper'|'note'|'milestone'), title, details, logged_by, timestamp.",
            "Attendance: id, child_id, date, status ('present'|'absent'|'late'), check_in_time, check_out_time, marked_by."
        ]),
        ("Communications & Learning", COLOR_CYAN, [
            "Message: id, child_id, sender_id, receiver_id, content, is_alert, alert_severity, timestamp, is_read.",
            "Notification: id, user_id, title, message, link, is_urgent, is_read.",
            "LearningResource: id, title, type, age_group, description, content_url_or_text."
        ])
    ]

    for i, (title, col, items) in enumerate(entities):
        x = start_left + i * (col_w + Inches(0.4))
        card = add_card(slide10, x, Inches(1.8), col_w, Inches(5.0))
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{title}\n"
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = col

        for it in items:
            pi = tf.add_paragraph()
            pi.text = f"• {it}\n"
            pi.font.size = Pt(12)
            pi.font.color.rgb = COLOR_TEXT_LIGHT

    slide10.notes_slide.notes_text_frame.text = (
        "The relational schema is engineered with SQLAlchemy ORM. Foreign key constraints bind children to both their guardian and assigned teacher. "
        "Cascading rules ensure orphaned records are cleaned if a child record is un-enrolled. Indexes on timestamp and date columns ensure fast queries."
    )

    # -------------------------------------------------------------
    # SLIDE 11: Security & OWASP Compliance
    # -------------------------------------------------------------
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11, COLOR_BG_DARK)
    add_header(slide11, "Security Architecture & Privacy Enforcement", "10. SECURITY & COMPLIANCE")

    sec_cards = [
        ("🔐 Object-Level Tenant Isolation", "Enforced via 'get_authorized_child()'. Every request strictly verifies whether the authenticated user owns or is assigned to the requested child ID. Unauthorized attempts trigger HTTP 403 Forbidden.", COLOR_ROSE, Inches(0.8), Inches(1.8)),
        ("🛡️ Gatekeeper Decorators", "Custom '@role_required(*roles)' decorator inspects current_user credentials. Caregiver accounts are permanently blocked from accessing child data until verified and approved by the Director.", COLOR_AMBER, Inches(6.9), Inches(1.8)),
        ("🔑 Cryptographic Credential Hashing", "Zero plaintext passwords. Uses PBKDF2/SHA-256 via Werkzeug's 'generate_password_hash' and 'check_password_hash' with salt protection against rainbow table attacks.", COLOR_ACCENT, Inches(0.8), Inches(4.5)),
        ("📁 Secure File Upload Sanitization", "Uploaded photos are renamed with random UUIDv4 strings and sanitized via 'secure_filename'. Restricts extensions strictly to valid image formats to prevent remote code execution (RCE).", COLOR_CYAN, Inches(6.9), Inches(4.5))
    ]

    for title, desc, col, l, t in sec_cards:
        card = add_card(slide11, l, t, Inches(5.6), Inches(2.35))
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_TEXT_LIGHT

    slide11.notes_slide.notes_text_frame.text = (
        "Child data requires high privacy standards. We mitigate top OWASP risks: Broken Access Control is prevented via "
        "object-level authorization in 'get_authorized_child'. Injection is mitigated via parameterized SQLAlchemy ORM queries. "
        "Insecure file uploads are safeguarded using UUID randomization and extension whitelisting."
    )

    # -------------------------------------------------------------
    # SLIDE 12: Automated Testing & Verification
    # -------------------------------------------------------------
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12, COLOR_BG_DARK)
    add_header(slide12, "Verification & Automated Test Suite (8/8 Suites)", "11. QUALITY ASSURANCE")

    c_left = add_card(slide12, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf_l = c_left.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "🧪 Test Execution (test_daycare.py)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT

    tests = [
        "1. Landing & Role Redirect: Verifies authentication routing for admin, caregiver, parent.",
        "2. Caregiver Approval Flow: Tests blocker on unapproved caregivers and Director approval.",
        "3. Activity Logging Sync: Validates educator log persistence and parent timeline visibility.",
        "4. Strict Child Data Isolation: Asserts 403 Forbidden when parents attempt to access unauthorized child IDs."
    ]
    for t in tests:
        pi = tf_l.add_paragraph()
        pi.text = f"\n✔ {t}"
        pi.font.size = Pt(12)
        pi.font.color.rgb = COLOR_TEXT_LIGHT

    c_right = add_card(slide12, Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.0))
    tf_r = c_right.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "🔬 Advanced Flow Verification"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    tests_adv = [
        "5. Emergency Alert Dispatch: Verifies high-priority alert generation and banner rendering.",
        "6. Daily Attendance Roster: Tests 1-click check-in and check-out state transition and roster updates.",
        "7. Learning Module Inspection: Verifies age-filtered curriculum rendering and reading mode.",
        "8. Password Reset Flow: Validates secure reset token generation and credential updates."
    ]
    for t in tests_adv:
        pi = tf_r.add_paragraph()
        pi.text = f"\n✔ {t}"
        pi.font.size = Pt(12)
        pi.font.color.rgb = COLOR_TEXT_LIGHT

    slide12.notes_slide.notes_text_frame.text = (
        "We backed our implementation with a comprehensive automated test suite in 'test_daycare.py'. "
        "Using Python unittest and an in-memory SQLite database, all 8 test suites pass with 100% success rate, "
        "validating role redirections, pending approval barriers, daily attendance marking, and security isolation."
    )

    # -------------------------------------------------------------
    # SLIDE 13: Deployment & Production Readiness
    # -------------------------------------------------------------
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13, COLOR_BG_DARK)
    add_header(slide13, "Deployment Architecture & Production Engineering", "12. DEPLOYMENT & DEVOPS")

    deploy_boxes = [
        ("🚀 WSGI Server / Gunicorn", "Configured with 'wsgi.py' and Gunicorn runner for multi-worker concurrency on Linux and container environments.", COLOR_PRIMARY),
        ("☁️ Cloud PaaS Ready", "Includes cloud-native 'Procfile' for 1-click deployment on platforms like Render, Railway, AWS ECS, or Heroku.", COLOR_ACCENT),
        ("🔒 12-Factor App Config", "Complete segregation of config and secrets via '.env' and 'python-dotenv'. Zero hardcoded credentials.", COLOR_AMBER),
        ("💾 Database Portability", "Zero-lock-in architecture: seamlessly toggle between SQLite for local dev/testing and PostgreSQL for production.", COLOR_CYAN)
    ]

    for i, (title, desc, col) in enumerate(deploy_boxes):
        col_idx = i % 2
        row_idx = i // 2
        x = Inches(0.8) + col_idx * Inches(6.1)
        y = Inches(1.8) + row_idx * Inches(2.45)
        card = add_card(slide13, x, y, Inches(5.6), Inches(2.25))
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_TEXT_LIGHT

    slide13.notes_slide.notes_text_frame.text = (
        "The project adheres to 12-Factor App methodology. It includes a WSGI entry point, a Procfile for cloud hosting, "
        "and clean separation of settings in .env. It supports both local zero-configuration SQLite and high-throughput PostgreSQL."
    )

    # -------------------------------------------------------------
    # SLIDE 14: Future Scope & Roadmap
    # -------------------------------------------------------------
    slide14 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide14, COLOR_BG_DARK)
    add_header(slide14, "Future Roadmap & Technical Extensions", "13. FUTURE SCOPE")

    scope_cards = [
        ("🤖 AI Developmental Analytics", "Train lightweight machine learning models to analyze milestone patterns, detect developmental speech/motor delays, and recommend tailored sensory exercises.", COLOR_CYAN, Inches(0.8)),
        ("📱 Native Mobile Apps", "Wrap the REST API with React Native or Flutter for native iOS and Android experiences with Apple Health & Google Fit biometric integration.", COLOR_ACCENT, Inches(4.75)),
        ("🚪 IoT Smart Gate & Access Control", "Connect smart RFID cards and biometric hardware directly to physical turnstiles at daycare entrance gates for automated, secure door unlock.", COLOR_PRIMARY, Inches(8.7))
    ]

    for title, desc, col, x in scope_cards:
        card = add_card(slide14, x, Inches(1.8), Inches(3.8), Inches(5.0))
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{title}\n"
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}\n\n"
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_TEXT_LIGHT

        p3 = tf.add_paragraph()
        p3.text = "Key Upgrades:\n• Stripe / Razorpay Tuition Billing\n• Real-Time Bus GPS Tracking\n• Multi-Language Audio Localization"
        p3.font.size = Pt(12)
        p3.font.color.rgb = COLOR_TEXT_MUTED

    slide14.notes_slide.notes_text_frame.text = (
        "For future enhancements, we have identified three major vectors: AI-driven developmental milestone analysis, "
        "native mobile applications with push notifications, and IoT smart-gate hardware integration for physical door access control."
    )

    # -------------------------------------------------------------
    # SLIDE 15: Conclusion & Q&A
    # -------------------------------------------------------------
    slide15 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide15, COLOR_BG_DARK)

    end_card = add_card(slide15, Inches(1.0), Inches(1.0), Inches(11.33), Inches(5.5))
    tf_end = end_card.text_frame
    tf_end.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_end.word_wrap = True

    p = tf_end.paragraphs[0]
    p.text = "🎓 Conclusion & Viva Defense"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_LIGHT

    p2 = tf_end.add_paragraph()
    p2.text = "Virtual Daycare Platform: Bridging Child Care with Real-Time Engineering"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(18)
    p2.font.color.rgb = COLOR_ACCENT

    p3 = tf_end.add_paragraph()
    p3.text = "\n• 100% Passing Automated Test Suite  • Production Ready WSGI Architecture\n• Zero-Leak Multi-Tenant Isolation  • Streamlined Daily Attendance Roster"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_TEXT_MUTED

    p4 = tf_end.add_paragraph()
    p4.text = "\n\nThank you! We welcome questions from the review committee."
    p4.alignment = PP_ALIGN.CENTER
    p4.font.size = Pt(16)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_PRIMARY

    slide15.notes_slide.notes_text_frame.text = (
        "In conclusion, the Virtual Daycare Platform delivers a robust, secure, and production-tested solution for "
        "modern early childhood institutions. It solves parental anxiety, modernizes manual tracking with an intuitive daily roster, "
        "and adheres to strict privacy isolation standards. Thank you, and I am now ready for questions."
    )

    output_path = "Virtual_Daycare_College_Presentation.pptx"
    prs.save(output_path)
    print(f"[SUCCESS] Presentation generated successfully at: {output_path}")

if __name__ == '__main__':
    create_presentation()
